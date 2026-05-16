"""Data Center — drag-drop Excel loading and data preview."""

from __future__ import annotations

import os
import shutil
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from PySide6.QtCore import Qt
from PySide6.QtGui import QDragEnterEvent, QDragLeaveEvent, QDropEvent
from PySide6.QtWidgets import (
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from qfluentwidgets import (
    BodyLabel,
    FluentIcon,
    InfoBar,
    MessageBox,
    PrimaryPushButton,
    PushButton,
    TitleLabel,
)


class DataPreviewWidget(QWidget):
    """Table preview of loaded data with drag-and-drop support."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setAcceptDrops(True)
        self._full_df: pd.DataFrame | None = None
        self._loading = False
        self._setup_ui()

    def _setup_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)

        header_row = QHBoxLayout()
        self.file_label = TitleLabel("未加载数据 (可将 Excel 拖拽至此)")
        self.file_label.setStyleSheet("font-size: 15px;")
        header_row.addWidget(self.file_label)

        self.browse_btn = PushButton("浏览文件")
        self.browse_btn.setIcon(FluentIcon.FOLDER)
        self.browse_btn.clicked.connect(self._browse_file)
        header_row.addWidget(self.browse_btn)

        self.sync_btn = PushButton("同步修改")
        self.sync_btn.setIcon(FluentIcon.SYNC)
        self.sync_btn.clicked.connect(self._sync_file)
        self.sync_btn.hide()  # Hidden until data is loaded
        header_row.addWidget(self.sync_btn)

        self._smart_btn_container = QWidget()
        self._smart_btn_layout = QHBoxLayout(self._smart_btn_container)
        self._smart_btn_layout.setContentsMargins(0, 0, 0, 0)
        self._smart_btn_layout.setSpacing(6)
        self._smart_btn_container.hide()
        header_row.addWidget(self._smart_btn_container)

        header_row.addStretch()

        self.time_label = BodyLabel("")
        self.time_label.setStyleSheet("color: #0078D4; margin-right: 15px;")
        header_row.addWidget(self.time_label)

        self.row_label = BodyLabel("")
        self.row_label.setStyleSheet("color: #666;")
        header_row.addWidget(self.row_label)
        layout.addLayout(header_row)

        from PySide6.QtWidgets import QFrame, QStackedWidget

        self.content_stack = QStackedWidget()

        # Tips Area (Visible by default)
        self.tips_container = QFrame()
        self.tips_container.setObjectName("tips_container")
        tips_layout = QVBoxLayout(self.tips_container)
        tips_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        tips_layout.setContentsMargins(40, 40, 40, 40)
        tips_layout.setSpacing(15)
        self.tips_container.setStyleSheet("""
            QFrame#tips_container {
                background-color: #ffffff;
                border: 2px dashed #ccc;
                border-radius: 12px;
            }
        """)

        welcome_title = TitleLabel("欢迎使用 Q-Caliper 数据中心")
        welcome_title.setStyleSheet("font-size: 18px; color: #0078D4; border: none; background: transparent;")
        tips_layout.addWidget(welcome_title, 0, Qt.AlignmentFlag.AlignCenter)
        tips_layout.addSpacing(10)

        # Inner container for left-aligned list within a centered area
        inner_tips_widget = QWidget()
        inner_tips_layout = QVBoxLayout(inner_tips_widget)
        inner_tips_layout.setContentsMargins(0, 0, 0, 0)
        inner_tips_layout.setSpacing(12)

        tips_data = [
            ("📁", "拖入数据文件可自动整理数据并在当前文件夹建立 [Q]_ 副本，支持双向实时同步。"),
            ("🔍", "支持前 20 行自动搜寻表头，并能自动识别合并复杂的双行表头。"),
            ("📝", "针对 Excel 合并单元格情况，系统将自动进行重复填充以确保数据完整性。"),
            ("📑", "支持多页表格（Sheet），检测到多个数据页面时将提示您进行选择。"),
            ("💡", "系统将智能推断数据角色并推荐分析模式，您也可以双击预览表头手动修正。"),
        ]

        for icon, text in tips_data:
            tip_item = QHBoxLayout()
            tip_item.setSpacing(12)

            icon_label = BodyLabel(icon)
            icon_label.setFixedWidth(25)  # Fixed width to align icons
            icon_label.setStyleSheet("font-size: 16px; background: transparent; border: none;")
            tip_item.addWidget(icon_label)

            text_label = BodyLabel(text)
            text_label.setStyleSheet("color: #444; font-size: 14px; background: transparent; border: none;")
            tip_item.addWidget(text_label)

            inner_tips_layout.addLayout(tip_item)

        tips_layout.addWidget(inner_tips_widget, 0, Qt.AlignmentFlag.AlignCenter)

        self.content_stack.addWidget(self.tips_container)

        self.table = QTableWidget()
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.table.verticalHeader().setDefaultSectionSize(28)
        self.table.setStyleSheet("""
            QTableWidget {
                border: 1px solid #edebe9;
                border-radius: 6px;
                background: white;
                alternate-background-color: #f8f9fa;
                gridline-color: #eee;
                font-size: 12px;
            }
            QTableWidget::item:selected {
                background: #cce4ff;
                color: black;
            }
            QHeaderView::section {
                background: #f0f2f5;
                border: none;
                border-bottom: 2px solid #0078D4;
                padding: 6px;
                font-weight: bold;
                font-size: 12px;
            }
        """)

        # 信号连接
        self.table.horizontalHeader().sectionDoubleClicked.connect(self._edit_header)
        self.table.itemChanged.connect(self._on_item_changed)

        self.content_stack.addWidget(self.table)
        self.content_stack.setCurrentWidget(self.tips_container)
        layout.addWidget(self.content_stack, 1)

    def set_sync_time(self, time_dt: datetime | None = None) -> None:
        if time_dt:
            self.time_label.setText(f"同步时间: {time_dt.strftime('%H:%M:%S')}")
        else:
            self.time_label.setText("")

    def dragEnterEvent(self, event: QDragEnterEvent) -> None:
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            target = self.table if self.content_stack.currentWidget() == self.table else self.tips_container
            target.setStyleSheet(
                target.styleSheet()
                .replace("border: 2px dashed #ccc;", "border: 2px dashed #0078D4; background-color: #f0f8ff;")
                .replace("border: 1px solid #edebe9;", "border: 2px dashed #0078D4; background: #f0f8ff;")
            )

    def dragLeaveEvent(self, event: QDragLeaveEvent) -> None:
        target = self.table if self.content_stack.currentWidget() == self.table else self.tips_container
        if self.content_stack.currentWidget() == self.table:
            target.setStyleSheet(
                target.styleSheet().replace(
                    "border: 2px dashed #0078D4; background: #f0f8ff;", "border: 1px solid #edebe9;"
                )
            )
        else:
            target.setStyleSheet(
                target.styleSheet().replace(
                    "border: 2px dashed #0078D4; background-color: #f0f8ff;",
                    "border: 2px dashed #ccc; background-color: #ffffff;",
                )
            )

    def dropEvent(self, event: QDropEvent) -> None:
        # Reset style
        self.dragLeaveEvent(None)  # type: ignore

        urls = event.mimeData().urls()
        if urls:
            path = urls[0].toLocalFile()
            self._load_file(path)

    def _browse_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择数据文件",
            "",
            "Excel 文件 (*.xlsx *.xlsm);;CSV 文件 (*.csv);;所有文件 (*)",
        )
        if path:
            self._load_file(path)

    def _load_file(self, path: str) -> None:
        parent = self.parent()
        while parent and not hasattr(parent, "load_file"):
            parent = parent.parent()
        if parent is not None and hasattr(parent, "load_file"):
            parent.load_file(path)

    def _edit_header(self, logical_index: int) -> None:
        if self.table.columnCount() == 0:
            return

        item = self.table.horizontalHeaderItem(logical_index)
        current_role = item.text() if item else "未分类"

        roles = ["测量值", "操作者", "零件", "时间", "因子", "未分类"]
        new_role, ok = QInputDialog.getItem(
            self,
            "修改数据角色",
            "请选择该列的数据角色:",
            roles,
            current=roles.index(current_role) if current_role in roles else 5,
            editable=False,
        )

        if ok and new_role:
            if item:
                item.setText(new_role)
            else:
                self.table.setHorizontalHeaderItem(logical_index, QTableWidgetItem(new_role))

            # 修改角色只同步状态，静默处理
            self._sync_file(silent=True)
            self._update_smart_btn()
            InfoBar.success("角色已更新", f"列 {logical_index} 的角色已更新为 '{new_role}'", parent=self, duration=2000)

    def _on_item_changed(self, item: QTableWidgetItem) -> None:
        if self._loading or self._full_df is None:
            return

        row = item.row()
        col = item.column()
        val = item.text()

        # 信号连接暂时关闭信号避免递归
        self.table.blockSignals(True)

        try:
            if row == 0:
                # 修改列名
                new_cols = list(self._full_df.columns)
                new_cols[col] = val
                self._full_df.columns = new_cols

                # 加粗显示列名
                font = item.font()
                font.setBold(True)
                item.setFont(font)
            else:
                # 修改数据 (行索引在表格中 offset 了 1)
                # 列上下文感知的类型转换
                converted_val: Any
                if not val:
                    converted_val = None
                else:
                    col_dtype = self._full_df.iloc[:, col].dtype
                    if col_dtype is object:
                        converted_val = val
                    else:
                        try:
                            converted_val = float(val) if "." in val or "e" in val.lower() else int(val)
                        except ValueError:
                            converted_val = val

                self._full_df.iloc[row - 1, col] = converted_val

            # 自动保存应该静默进行，避免频繁弹窗干扰用户
            self._sync_file(silent=True)
        except Exception as e:
            InfoBar.error("数据更新失败", str(e), parent=self, duration=-1)
        finally:
            self.table.blockSignals(False)

    def _sync_file(self, silent: bool = False) -> None:
        if self._full_df is None:
            return

        parent = self.parent()
        while parent and not hasattr(parent, "sync_file"):
            parent = parent.parent()
        if parent is not None and hasattr(parent, "sync_file"):
            parent.sync_file(self._full_df, silent=silent)

    def _get_current_dataframe(self) -> pd.DataFrame:
        """Constructs a DataFrame from the current QTableWidget contents. Row 0 contains column names.
        DEPRECATED: Now uses self._full_df directly to avoid truncation.
        """
        return self._full_df if self._full_df is not None else pd.DataFrame()

    def _get_recommendations(self) -> list[tuple[str, str, str]]:
        """Return list of (panel_attr, label, icon_name) recommendations."""
        mapping = self.get_role_mapping()
        recs: list[tuple[str, str, str]] = []
        if all(r in mapping for r in ["测量值", "操作者", "零件"]):
            recs.append(("grr_panel", "GRR", "MARKET"))
        if "测量值" in mapping:
            recs.append(("cpk_panel", "CPK", "CARE_RIGHT_SOLID"))
            recs.append(("spc_panel", "SPC", "CHART"))
        if "因子" in mapping:
            recs.append(("doe_panel", "DOE", "CODE"))
        return recs

    def _navigate_to(self, panel_attr: str) -> None:
        """Navigate to a specific analysis panel."""
        mapping = self.get_role_mapping()
        main_win = self.window()

        # Handle multiple measurement columns
        if "测量值" in mapping and len(mapping["测量值"]) > 1:
            col, ok = QInputDialog.getItem(
                self,
                "选择分析列",
                "检测到多个测量值列，请选择:",
                mapping["测量值"],
                0,
                False,
            )
            if ok:
                mapping["测量值"] = [col]
            else:
                return

        # Special handling for merged MSA/GRR panel
        if panel_attr == "grr_panel":
            panel_attr = "msa_panel"

        target_panel = getattr(main_win, panel_attr, None)
        if target_panel and hasattr(main_win, "switchTo"):
            if hasattr(target_panel, "set_selected_columns"):
                target_panel.set_selected_columns(mapping)

            # If navigating to MSA for GRR, ensure the GRR tab is selected
            if panel_attr == "msa_panel" and hasattr(target_panel, "pivot"):
                target_panel.pivot.setCurrentItem("grr")
                if hasattr(target_panel, "stacked_widget") and hasattr(target_panel, "grr_interface"):
                    target_panel.stacked_widget.setCurrentWidget(target_panel.grr_interface)

            main_win.switchTo(target_panel)
        else:
            # Only show warning if navigation target is fundamentally missing or roles aren't mapped
            recs = self._get_recommendations()
            if not recs:
                InfoBar.warning(
                    "分析推荐", "请先通过双击表头为列分配角色（如：测量值、零件、操作者）", parent=self, duration=-1
                )
            else:
                InfoBar.error("系统错误", f"无法定位分析模块: {panel_attr}", parent=self, duration=-1)

    def get_role_mapping(self) -> dict[str, list[str]]:
        """Extract mapping of roles to column names from the table."""
        mapping: dict[str, list[str]] = {}
        for j in range(self.table.columnCount()):
            header_item = self.table.horizontalHeaderItem(j)
            if not header_item:
                continue
            role = header_item.text()
            # The actual column name is in the 0-th row of the table
            col_item = self.table.item(0, j)
            if not col_item:
                continue
            col_name = col_item.text()
            if role not in mapping:
                mapping[role] = []
            mapping[role].append(col_name)
        return mapping

    def _update_smart_btn(self) -> None:
        """Rebuild smart recommendation buttons based on current roles."""
        # Clear existing buttons
        while self._smart_btn_layout.count() > 0:
            item = self._smart_btn_layout.takeAt(0)
            if item is not None:
                w = item.widget()
                if w is not None:
                    w.deleteLater()

        recs = self._get_recommendations()
        if not recs:
            btn = PrimaryPushButton("分析推荐")
            btn.setIcon(FluentIcon.INFO)
            btn.clicked.connect(
                lambda: InfoBar.warning(
                    "分析推荐",
                    "请先通过双击表头为列分配角色（如：测量值、零件、操作者）",
                    parent=self,
                )
            )
            self._smart_btn_layout.addWidget(btn)
        else:
            for panel_attr, label, icon_name in recs:
                btn = PrimaryPushButton(label)
                icon = getattr(FluentIcon, icon_name, FluentIcon.CARE_RIGHT_SOLID)
                btn.setIcon(icon)
                btn.clicked.connect(lambda checked=False, p=panel_attr: self._navigate_to(p))
                self._smart_btn_layout.addWidget(btn)

        self._smart_btn_container.show()

    def load_dataframe(self, df: pd.DataFrame, filename: str) -> None:
        self._loading = True
        self.table.blockSignals(True)
        self._full_df = df.copy()  # 保存全量数据副本

        self.file_label.setText(filename)
        n_rows, n_cols = df.shape
        self.row_label.setText(f"{n_rows} 行 x {n_cols} 列")

        self.table.clear()

        # 智能推断数据角色
        roles = []
        for col in df.columns:
            col_str = str(col).lower()
            if any(k in col_str for k in ("人", "操作", "检验员", "operator", "appraiser")):
                roles.append("操作者")
            elif any(k in col_str for k in ("零件", "产品", "part", "item", "产品号")):
                roles.append("零件")
            elif any(k in col_str for k in ("测量", "值", "value", "data", "数据", "结果")):
                roles.append("测量值")
            elif any(k in col_str for k in ("日期", "时间", "date", "time")):
                roles.append("时间")
            elif any(k in col_str for k in ("因子", "factor", "条件", "水平", "level", "正交")):
                roles.append("因子")
            elif pd.api.types.is_float_dtype(df[col]):
                roles.append("测量值")
            else:
                roles.append("未分类")

        max_preview_rows = min(n_rows, 100)
        self.table.setRowCount(max_preview_rows + 1)
        self.table.setColumnCount(n_cols)
        self.table.setHorizontalHeaderLabels(roles)

        # 第 0 行放置真实表头
        for j in range(n_cols):
            item = QTableWidgetItem(str(df.columns[j]))
            font = item.font()
            font.setBold(True)
            item.setFont(font)
            item.setBackground(Qt.GlobalColor.lightGray)
            self.table.setItem(0, j, item)

        # 第 1 行起放置真实数据
        for i in range(max_preview_rows):
            for j in range(n_cols):
                val = df.iloc[i, j]
                item = QTableWidgetItem(str(val) if pd.notna(val) else "")
                self.table.setItem(i + 1, j, item)

        self.table.resizeColumnsToContents()

        # Switch visibility
        self.content_stack.setCurrentWidget(self.table)

        self.table.blockSignals(False)
        self._loading = False
        self.sync_btn.show()
        self._update_smart_btn()


class DataCenterWidget(QWidget):
    """Data Center page — file loading and preview."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setObjectName("data_center")
        self.df: pd.DataFrame | None = None
        self.filepath: str | None = None
        self._setup_ui()

    def _setup_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 10, 20, 20)

        header = TitleLabel("数据中心")
        header.setStyleSheet("font-size: 20px; font-weight: bold;")
        layout.addWidget(header)

        desc = BodyLabel("拖拽或选择数据文件, 自动预览并启用可用分析功能")
        desc.setStyleSheet("color: #666; margin-bottom: 8px;")
        layout.addWidget(desc)

        self.preview = DataPreviewWidget(self)
        layout.addWidget(self.preview, 1)

    def load_file(self, path: str) -> None:
        try:
            p = Path(path)
            suffix = p.suffix.lower()

            if suffix not in (".xlsx", ".xlsm", ".csv"):
                InfoBar.error(
                    "不支持的格式", f"不支持 {suffix} 格式, 请使用 .xlsx / .xlsm / .csv", parent=self, duration=-1
                )
                return

            # --- Read-only file handling ---
            if suffix != ".csv" and not os.access(path, os.W_OK):
                title = "文件为只读"
                content = f"文件 '{p.name}' 当前为只读状态，无法清洗和编辑。\n是否取消只读属性后继续？"
                w = MessageBox(title, content, self)
                w.yesButton.setText("取消只读并继续")
                w.cancelButton.setText("取消")
                if not w.exec():
                    return
                try:
                    import stat

                    p.chmod(p.stat().st_mode | stat.S_IWRITE)
                except Exception as e:
                    InfoBar.error("操作失败", f"无法修改文件权限: {e!s}", parent=self, duration=-1)
                    return

            target_path = p
            if not p.name.startswith("[Q]_"):
                target_name = f"[Q]_{p.name}"
                target_path = p.with_name(target_name)

                if target_path.exists():
                    title = "发现已存在的分析副本"
                    content = f"文件夹中已存在分析副本 '{target_path.name}'。\n是否用当前选中的源文件覆盖它？(注意：覆盖后将丢失之前在副本中的所有修改)"
                    w = MessageBox(title, content, self)
                    w.yesButton.setText("覆盖")
                    w.cancelButton.setText("取消")
                    if not w.exec():
                        return  # 用户取消

                # 无论是新生成还是确认覆盖，都进行拷贝
                shutil.copy2(path, target_path)

            # --- Multi-sheet detection ---
            sheet_name: str | int | None = None
            if suffix != ".csv":
                try:
                    from openpyxl import load_workbook

                    wb = load_workbook(target_path, read_only=True, data_only=True)
                    non_empty_sheets: list[str] = []
                    for sname in wb.sheetnames:
                        ws = wb[sname]
                        has_data = False
                        for row in ws.iter_rows(max_row=5, values_only=True):
                            if any(v is not None for v in row):
                                has_data = True
                                break
                        if has_data:
                            non_empty_sheets.append(sname)
                    wb.close()

                    if len(non_empty_sheets) > 1:
                        chosen, ok = QInputDialog.getItem(
                            self,
                            "选择数据表",
                            f"文件包含 {len(non_empty_sheets)} 个非空数据表，请选择要分析的表:",
                            non_empty_sheets,
                            0,
                            False,
                        )
                        if not ok:
                            return
                        sheet_name = chosen
                except Exception:
                    pass

            df = self._smart_read_and_clean(target_path, sheet_name=sheet_name)
            if df is None:
                return

            self.df = df
            self.filepath = str(target_path)
            self.preview.load_dataframe(df, target_path.name)
            self.preview.set_sync_time(datetime.now())

            # Record timestamp for future sync
            self._last_modified_time = os.path.getmtime(self.filepath)

            InfoBar.success(
                "加载成功",
                f"已加载 {target_path.name} ({len(df)} 行 x {len(df.columns)} 列)",
                parent=self,
                duration=2000,
            )

            main_win = self.window()
            if hasattr(main_win, "cpk_panel"):
                main_win.cpk_panel.set_dataframe(df, target_path.name)
            if hasattr(main_win, "grr_panel"):
                main_win.grr_panel.set_dataframe(df, target_path.name)
            if hasattr(main_win, "msa_panel"):
                main_win.msa_panel.set_dataframe(df, target_path.name)
            if hasattr(main_win, "spc_panel"):
                main_win.spc_panel.set_dataframe(df, target_path.name)
            if hasattr(main_win, "doe_panel"):
                main_win.doe_panel.set_dataframe(df, target_path.name)

        except Exception as e:
            InfoBar.error("加载失败", str(e), parent=self, duration=-1)

    def sync_file(self, current_df: pd.DataFrame, silent: bool = False) -> None:
        if not self.filepath:
            return

        try:
            current_mtime = os.path.getmtime(self.filepath)

            # Check for external modification (with 1 second tolerance)
            if hasattr(self, "_last_modified_time") and current_mtime > self._last_modified_time + 1:
                title = "检测到外部修改"
                content = "文件在外部被修改，是否重新加载外部内容？（选择否将用当前界面内容覆盖外部文件）"
                w = MessageBox(title, content, self)
                w.yesButton.setText("重新加载")
                w.cancelButton.setText("覆盖")
                if w.exec():
                    self.load_file(self.filepath)
                    return

            # Overwrite external file with current UI data
            suffix = Path(self.filepath).suffix.lower()
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", category=UserWarning)
                if suffix == ".csv":
                    current_df.to_csv(self.filepath, index=False)
                else:
                    current_df.to_excel(self.filepath, engine="openpyxl", index=False)

            self.df = current_df
            self._last_modified_time = os.path.getmtime(self.filepath)

            # Update other panels
            main_win = self.window()
            if hasattr(main_win, "cpk_panel"):
                main_win.cpk_panel.set_dataframe(current_df, Path(self.filepath).name)
            if hasattr(main_win, "grr_panel"):
                main_win.grr_panel.set_dataframe(current_df, Path(self.filepath).name)
            if hasattr(main_win, "msa_panel"):
                main_win.msa_panel.set_dataframe(current_df, Path(self.filepath).name)
            if hasattr(main_win, "spc_panel"):
                main_win.spc_panel.set_dataframe(current_df, Path(self.filepath).name)
            if hasattr(main_win, "doe_panel"):
                main_win.doe_panel.set_dataframe(current_df, Path(self.filepath).name)

            self.preview.set_sync_time(datetime.now())

            if not silent:
                InfoBar.success("同步成功", "界面修改已保存至文件并更新所有分析模块", parent=self, duration=2000)

        except Exception as e:
            InfoBar.error("同步失败", str(e), parent=self, duration=-1)

    @staticmethod
    def _fill_merged_cells(
        df: pd.DataFrame,
        merged_ranges: list[tuple[int, int, int, int]],
    ) -> pd.DataFrame:
        """Fill NaN cells inside merged ranges with the top-left value.

        Parameters
        ----------
        df : DataFrame already read (0-indexed, header=None)
        merged_ranges : list of (min_row, max_row, min_col, max_col) 1-indexed
        """
        nrows, ncols = df.shape
        for min_row, max_row, min_col, max_col in merged_ranges:
            if min_row > nrows or min_col > ncols:
                continue
            top_left = df.iloc[min_row - 1, min_col - 1]
            if pd.isna(top_left):
                continue
            for r in range(min_row - 1, min(max_row, nrows)):
                for c in range(min_col - 1, min(max_col, ncols)):
                    if pd.isna(df.iloc[r, c]):
                        df.iloc[r, c] = top_left
        return df

    @staticmethod
    def _merge_double_headers(row1: pd.Series, row2: pd.Series) -> list[str]:
        """Merge two header rows into a single row using underscore join."""
        result: list[str] = []
        row1_ffilled = row1.ffill()
        for i in range(len(row1_ffilled)):
            v1 = str(row1_ffilled.iloc[i]).strip() if pd.notna(row1_ffilled.iloc[i]) else ""
            v2 = str(row2.iloc[i]).strip() if pd.notna(row2.iloc[i]) else ""
            if v1 and v2 and v1 != v2:
                result.append(f"{v1}_{v2}")
            elif v1:
                result.append(v1)
            elif v2:
                result.append(v2)
            else:
                result.append(f"col_{i}")
        return result

    @staticmethod
    def _normalize_column_types(df: pd.DataFrame) -> pd.DataFrame:
        """Ensure object columns that look like pure numeric strings stay as strings."""
        for i in range(len(df.columns)):
            series = df.iloc[:, i]
            if series.dtype != object:
                continue
            non_null = series.dropna()
            if non_null.empty:
                continue
            all_numeric = non_null.apply(
                lambda x: (
                    isinstance(x, (int, float))
                    or (isinstance(x, str) and x.strip().replace(".", "", 1).replace("-", "", 1).isdigit())
                )
            )
            if all_numeric.all():
                df.iloc[:, i] = series.astype(str)
        return df

    def _smart_read_and_clean(
        self,
        file_path: Path,
        sheet_name: str | int | None = None,
    ) -> pd.DataFrame | None:
        """Smartly detects header, handles merged cells & double headers, drops all-NaN rows/cols."""
        try:
            suffix = file_path.suffix.lower()
            read_kw: dict[str, Any] = {}
            if sheet_name is not None:
                read_kw["sheet_name"] = sheet_name

            if suffix == ".csv":
                raw_df = pd.read_csv(file_path, header=None, nrows=20)
            else:
                raw_df = pd.read_excel(
                    file_path,
                    engine="openpyxl",
                    header=None,
                    nrows=20,
                    **read_kw,
                )

            # --- Merged cells (Excel only) ---
            merged_ranges: list[tuple[int, int, int, int]] = []
            if suffix != ".csv":
                try:
                    from openpyxl import load_workbook

                    wb = load_workbook(file_path, data_only=True)
                    if sheet_name is not None and sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                    elif sheet_name is not None and isinstance(sheet_name, int):
                        ws = wb.worksheets[sheet_name]
                    else:
                        ws = wb.active
                    if ws is not None:
                        for mr in ws.merged_cells.ranges:
                            merged_ranges.append(
                                (mr.min_row, mr.max_row, mr.min_col, mr.max_col),
                            )
                    wb.close()
                except Exception:
                    pass

                if merged_ranges:
                    raw_df = self._fill_merged_cells(raw_df, merged_ranges)

            # --- Find best header row ---
            best_header_idx = 0
            max_non_nulls = 0
            for i in range(len(raw_df)):
                non_null_count = raw_df.iloc[i].count()
                if non_null_count > max_non_nulls:
                    max_non_nulls = non_null_count
                    best_header_idx = i

            if max_non_nulls == 0:
                InfoBar.warning("未能识别有效数据", "在前 20 行中没有找到任何有效的表头或数据。", parent=self)
                return None

            # --- Double header detection (check adjacent row only) ---
            merged_headers: list[str] | None = None
            data_start: int | None = None

            if best_header_idx + 1 < len(raw_df):
                adj_row = raw_df.iloc[best_header_idx + 1]
                adj_vals = adj_row.dropna()
                adj_is_text = all(not isinstance(v, (int, float)) for v in adj_vals)
                adj_count = len(adj_vals)
                if adj_is_text and adj_count >= 2 and adj_count >= max_non_nulls * 0.3:
                    merged_headers = self._merge_double_headers(
                        raw_df.iloc[best_header_idx],
                        adj_row,
                    )
                    data_start = best_header_idx + 2

            # --- Read full dataframe ---
            if suffix == ".csv":
                df = pd.read_csv(file_path, header=best_header_idx)
            else:
                if merged_headers is not None and data_start is not None:
                    df = pd.read_excel(
                        file_path,
                        engine="openpyxl",
                        header=None,
                        skiprows=data_start,
                        **read_kw,
                    )
                    df.columns = merged_headers[: len(df.columns)]
                else:
                    df = pd.read_excel(
                        file_path,
                        engine="openpyxl",
                        header=best_header_idx,
                        **read_kw,
                    )

            # Fill merged cells in the full data body
            if merged_ranges and suffix != ".csv":
                row_offset = data_start if data_start is not None else best_header_idx + 1
                adjusted: list[tuple[int, int, int, int]] = []
                for mr in merged_ranges:
                    adj_min = mr[0] - 1 - row_offset
                    adj_max = mr[1] - row_offset
                    if adj_max < 0:
                        continue
                    adjusted.append((max(adj_min, 0) + 1, adj_max + 1, mr[2], mr[3]))
                df = df.reset_index(drop=True)
                if adjusted:
                    df = self._fill_merged_cells(df, adjusted)

            # Clean empty rows and columns
            df = df.dropna(how="all", axis=0)
            df = df.dropna(how="all", axis=1)

            # Normalize column types
            df = self._normalize_column_types(df)

            # Save the cleaned dataframe back to the [Q]_ file
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", category=UserWarning)
                if suffix == ".csv":
                    df.to_csv(file_path, index=False)
                else:
                    df.to_excel(file_path, engine="openpyxl", index=False)

            return df
        except Exception as e:
            InfoBar.error("数据清洗失败", f"无法自动识别表头或清洗数据: {e!s}", parent=self, duration=-1)
            return None
